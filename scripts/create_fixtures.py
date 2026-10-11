from pathlib import Path
import zipfile
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from docx import Document
root=Path(__file__).resolve().parents[1]/'tests'/'fixtures';root.mkdir(parents=True,exist_ok=True)
c=canvas.Canvas(str(root/'Learning.pdf'),pagesize=A4)
for page in (1,2):
 c.setFont('Helvetica',22);c.drawString(55,770,'Learning that lasts '+str(page))
 c.setFont('Helvetica',14);c.drawString(55,715,'A resilient learner returns to a difficult idea.')
 c.drawString(55,685,'Retrieval practice can improve long-term retention.')
 c.showPage()
c.save()
d=Document();d.add_heading('Learning with context',0);d.add_paragraph('A resilient learner returns to a challenging problem.');d.add_heading('Retrieval practice',1);d.add_paragraph('Before you reveal the answer, try to remember its meaning.');d.add_paragraph('保留原句，使用主动回忆和间隔复习。');table=d.add_table(rows=2,cols=2);table.cell(0,0).text='Word';table.cell(0,1).text='Meaning';table.cell(1,0).text='resilient';table.cell(1,1).text='有韧性的';d.save(root/'Learning.docx')
with zipfile.ZipFile(root/'Learning.epub','w') as z:
 z.writestr('mimetype','application/epub+zip',compress_type=zipfile.ZIP_STORED)
 z.writestr('META-INF/container.xml','<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')
 z.writestr('OEBPS/content.opf','<?xml version="1.0"?><package version="3.0" unique-identifier="id" xmlns="http://www.idpf.org/2007/opf"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="id">memreader-test</dc:identifier><dc:title>Learning EPUB</dc:title><dc:language>en</dc:language><meta property="dcterms:modified">2026-10-10T00:00:00Z</meta></metadata><manifest><item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/><item id="c1" href="c1.xhtml" media-type="application/xhtml+xml"/><item id="c2" href="c2.xhtml" media-type="application/xhtml+xml"/></manifest><spine><itemref idref="c1"/><itemref idref="c2"/></spine></package>')
 z.writestr('OEBPS/nav.xhtml','<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops"><head><title>Contents</title></head><body><nav epub:type="toc"><ol><li><a href="c1.xhtml">First encounter</a></li><li><a href="c2.xhtml">Lasting memory</a></li></ol></nav></body></html>')
 for i in (1,2):z.writestr('OEBPS/c'+str(i)+'.xhtml','<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Chapter '+str(i)+'</title></head><body><h1>Chapter '+str(i)+'</h1><p>A resilient learner remembers ideas in context.</p><p>读书时保留原句，再通过主动回忆复习。</p></body></html>')
print('PDF, DOCX, EPUB fixtures ready')
