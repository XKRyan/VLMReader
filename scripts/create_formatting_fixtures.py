from pathlib import Path
from io import BytesIO
import zipfile
from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Inches, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor

root = Path(__file__).resolve().parents[1] / 'tests' / 'fixtures'
root.mkdir(exist_ok=True)
image = Image.new('RGB', (480, 140), '#f1f4fc')
draw = ImageDraw.Draw(image)
for x, color, label in [(40, '#5865da', 'READ'), (195, '#9477cf', 'RECALL'), (350, '#32a58c', 'REVIEW')]:
    draw.rounded_rectangle((x, 35, x+90, 105), radius=12, fill=color)
    draw.text((x+16, 61), label, fill='white')
    if x < 350:
        draw.line((x+100, 70, x+145, 70), fill='#98a0b8', width=4)
image.save(root / 'formatting-figure.png')

doc = Document()
doc.add_heading('Document formatting', 0)
p = doc.add_paragraph('A resilient learner keeps the original ')
r = p.add_run('yellow highlight'); r.font.highlight_color = WD_COLOR_INDEX.YELLOW
p.add_run(' and ')
r = p.add_run('red text'); r.font.color.rgb = RGBColor.from_string('C00000'); r.bold = True
p.add_run('.')
style = doc.styles.add_style('Reading accent', WD_STYLE_TYPE.PARAGRAPH)
style.font.color.rgb = RGBColor.from_string('0066CC')
doc.add_paragraph('Paragraph style carries blue ink.', style)
char = doc.styles.add_style('Marked term', WD_STYLE_TYPE.CHARACTER)
char.font.color.rgb = RGBColor.from_string('800000'); char.font.highlight_color = WD_COLOR_INDEX.TURQUOISE; char.font.bold = True
p = doc.add_paragraph('Active '); p.add_run('retrieval', char); p.add_run(' helps memory.')
p = doc.add_paragraph()
r = p.add_run('Dark blue highlight'); r.font.highlight_color = WD_COLOR_INDEX.DARK_BLUE
p.add_run(' / '); r = p.add_run('underlined detail'); r.underline = True
p = doc.add_paragraph()
r = p.add_run('Theme tint'); color = OxmlElement('w:color'); color.set(qn('w:val'), '95B3D7'); color.set(qn('w:themeColor'), 'accent1'); color.set(qn('w:themeTint'), '99'); r._element.get_or_add_rPr().append(color)
p.add_run(' / ')
r = p.add_run('Shaded text'); shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), 'D8EAD3'); r._element.get_or_add_rPr().append(shd)
doc.add_picture(str(root / 'formatting-figure.png'), width=Inches(4))
doc.add_paragraph('Figure 1. Read, recall and review.')
table = doc.add_table(rows=2, cols=2)
table.cell(0, 0).text = 'Method'; table.cell(0, 1).text = 'Example'
table.cell(1, 0).text = 'Retrieval'; r = table.cell(1, 1).paragraphs[0].add_run('Colored table cell'); r.font.color.rgb = RGBColor.from_string('008000'); r.font.highlight_color = WD_COLOR_INDEX.YELLOW
doc.save(root / 'Formatting.docx')

html = '''<!doctype html><html><head><style>
.original {background-color:#ffff00;color:#800000;font-family:Courier;font-size:90px;position:fixed}
.original em {color:#0066cc}
.app-sidebar {color:red}
</style></head><body><h1>HTML formatting</h1>
<p>Keep <span class="original">yellow highlighting <em>and blue words</em></span> in context.</p>
<p><font color="#008000">Legacy green ink</font> / <span style="color:rgb(192,0,0);background:#d8ead3">Inline red ink</span></p>
<img src="formatting-figure.png" alt="Read recall review" onerror="window.documentExecuted=true">
<p>Figure. A local image stays with the book.</p><script>window.documentExecuted=true</script>
<span style="background-image:url(https://invalid.example/tracker);font-size:200px;position:fixed;inset:0" data-ui>Ordinary text</span>
</body></html>'''
(root / 'Formatting.html').write_text(html, encoding='utf-8')
(root / 'Formatting.md').write_text('# Markdown formatting\n\nA **resilient** learner.\n\n![Read recall review](formatting-figure.png)\n\n<span style="background-color:yellow;color:#800000">Highlighted Markdown</span>\n', encoding='utf-8')

with zipfile.ZipFile(root / 'Formatting.epub', 'w') as z:
    z.writestr('mimetype', 'application/epub+zip', compress_type=zipfile.ZIP_STORED)
    z.writestr('META-INF/container.xml', '<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')
    z.writestr('OEBPS/content.opf', '<?xml version="1.0"?><package version="3.0" unique-identifier="id" xmlns="http://www.idpf.org/2007/opf"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="id">vlm-formatting-test</dc:identifier><dc:title>EPUB formatting</dc:title><dc:language>en</dc:language><meta property="dcterms:modified">2026-10-12T00:00:00Z</meta></metadata><manifest><item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/><item id="c1" href="c1.xhtml" media-type="application/xhtml+xml"/><item id="img" href="figure.png" media-type="image/png"/></manifest><spine><itemref idref="c1"/></spine></package>')
    z.writestr('OEBPS/nav.xhtml', '<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops"><head><title>Contents</title></head><body><nav epub:type="toc"><ol><li><a href="c1.xhtml">Formatting</a></li></ol></nav></body></html>')
    z.writestr('OEBPS/c1.xhtml', '<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Formatting</title><style>.source-highlight {background-color:#ffff00;color:#800000} img {max-width:100%}</style></head><body><h1>EPUB formatting</h1><p>A <span class="source-highlight">resilient learner</span> recalls ideas.</p><img src="figure.png" alt="Read recall review"/></body></html>')
    z.writestr('OEBPS/figure.png', (root / 'formatting-figure.png').read_bytes())

c = canvas.Canvas(str(root / 'Formatting.pdf'), pagesize=(595, 842))
c.setFont('Helvetica', 22); c.drawString(55, 780, 'PDF formatting')
c.setFillColor(HexColor('#ffff00')); c.rect(55, 700, 310, 26, fill=1, stroke=0)
c.setFillColor(HexColor('#800000')); c.setFont('Helvetica', 16); c.drawString(60, 708, 'Highlighted source text stays colored.')
c.drawImage(str(root / 'formatting-figure.png'), 55, 490, width=480, height=140)
c.save()
print('Formatting fixtures generated: DOCX, HTML, Markdown, EPUB and PDF.')
